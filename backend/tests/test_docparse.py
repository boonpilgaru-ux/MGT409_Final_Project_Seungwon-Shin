"""Hostile-document defenses in app/docparse.py: killable child process, PDF/DOCX caps."""

from __future__ import annotations

import io
import json
import subprocess
import sys
import time
import zipfile
import zlib

import docx
import pytest

from app import context, docparse
from app.models import UploadedFile
from tests.test_context import make_docx, make_pdf

_FONT = b"/Resources << /Font << /F1 << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> >> >>"


def pdf_pages(streams: list[bytes], page_streams: list[int]) -> bytes:
    """A PDF whose page i draws content stream ``page_streams[i]`` (pages may share a stream)."""
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    first_page, first_stream = 3, 3 + len(page_streams)
    kids = " ".join(f"{first_page + i} 0 R" for i in range(len(page_streams))).encode()
    objs.append(b"<< /Type /Pages /Kids [" + kids + b"] /Count %d >>" % len(page_streams))
    for s in page_streams:
        objs.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] " + _FONT
                    + b" /Contents %d 0 R >>" % (first_stream + s))
    for content in streams:
        comp = zlib.compress(content)
        objs.append(b"<< /Length %d /Filter /FlateDecode >>\nstream\n" % len(comp) + comp + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


def text_op(text: str) -> bytes:
    return f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET\n".encode("latin-1")


# --------------------------------------------------------------------------- PDF caps (in-process)


def test_pages_sharing_one_content_stream_are_read_once():
    pdf = pdf_pages([text_op("Zine maker")], [0] * 40)  # 40 pages -> the same stream
    assert docparse.pdf_text(pdf).count("Zine maker") == 1


def test_oversized_page_content_is_skipped_but_other_pages_are_read(monkeypatch):
    monkeypatch.setattr(docparse, "PDF_PAGE_MAX_CONTENT", 2_000)
    pdf = pdf_pages([b"0 0 m\n" * 1_000 + text_op("hidden"), text_op("Riso printer")], [0, 1])
    text = docparse.pdf_text(pdf)
    assert "Riso printer" in text and "hidden" not in text


def test_total_decoded_content_is_capped(monkeypatch):
    monkeypatch.setattr(docparse, "PDF_TOTAL_MAX_CONTENT", 100)
    pdf = pdf_pages([text_op(f"page {i} " + "x" * 30) for i in range(5)], list(range(5)))
    text = docparse.pdf_text(pdf)
    assert "page 0" in text and "page 4" not in text


def pdf_with_form(page_ops: bytes, form_ops: bytes, *, self_nested: bool = False) -> bytes:
    """Two pages: page 1 runs ``page_ops`` with form XObject /X0 (= ``form_ops``) in its resources;
    page 2 just says "Riso printer". ``self_nested`` makes /X0 list itself in its own resources."""
    form_res = b" /Resources << /XObject << /X0 6 0 R >> >>" if self_nested else b" " + _FONT
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 << /Type /Font "
        b"/Subtype /Type1 /BaseFont /Helvetica >> >> /XObject << /X0 6 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] " + _FONT + b" /Contents 7 0 R >>",
    ]
    for content, extra in ((page_ops, b""), (form_ops, b"/Type /XObject /Subtype /Form /BBox [0 0 612 792]" + form_res),
                           (text_op("Riso printer"), b"")):
        comp = zlib.compress(content)
        objs.append(b"<< " + extra + b" /Length %d /Filter /FlateDecode >>\nstream\n" % len(comp) + comp + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


_FORM = text_op("from the form") + b"0 0 m\n" * 200  # ~1.2 KB of operators


def test_form_xobjects_are_read_when_the_page_budget_allows(monkeypatch):
    monkeypatch.setattr(docparse, "PDF_PAGE_MAX_CONTENT", 2_000)
    text = docparse.pdf_text(pdf_with_form(b"/X0 Do\n", _FORM))
    assert "from the form" in text and "Riso printer" in text


@pytest.mark.parametrize(
    "page_ops",
    [
        b"/X0 Do\n" * 3,  # each invocation costs the form's full size again (3 x 1.2 KB > 2 KB)
        b"/X#30 Do\n" * 3,  # the same name, hex-escaped
        b"/X0 % hide it\nDo\n" * 3,  # a comment between operand and operator
        b"BT /F1 9 Tf (%) Tj ET /X0 Do\n" * 3,  # a "%" in a string: not a comment, so the Do still runs
        b"/X0[]Do\n" * 3,  # extra operands between the name and Do (pypdf paints operands[0])
        b"/X0()Do\n" * 3,
    ],
)
def test_painting_a_form_many_times_counts_toward_the_page_cap(monkeypatch, page_ops):
    monkeypatch.setattr(docparse, "PDF_PAGE_MAX_CONTENT", 2_000)
    text = docparse.pdf_text(pdf_with_form(page_ops, _FORM))
    assert "from the form" not in text and "Riso printer" in text  # page 1 skipped, page 2 read


def test_a_form_that_paints_itself_is_skipped_not_followed():
    started = time.perf_counter()
    text = docparse.pdf_text(pdf_with_form(b"/X0 Do\n", text_op("loop") + b"/X0 Do\n", self_nested=True))
    assert "Riso printer" in text and "loop" not in text
    assert time.perf_counter() - started < 2


async def test_isolated_parse_survives_the_form_xobject_bomb():
    """The verifier's residual: a tiny page saying "/X0 Do" over a form with MBs of "q " (pypdf
    keeps one graphics state per "q") used to reach ~1.2 GB in the child before the kill."""
    started = time.perf_counter()
    text = await docparse.extract_isolated("pdf", pdf_with_form(b"/X0 Do\n", b"q " * (3 * 1024 * 1024)), 20)
    assert "Riso printer" in text
    assert time.perf_counter() - started < 10


# --------------------------------------------------------------------------- DOCX caps (in-process)


def _rezip(data: bytes, extra: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item))
        for name, body in extra.items():
            dst.writestr(name, body)
    return buf.getvalue()


def test_docx_with_too_many_entries_is_refused():
    data = _rezip(make_docx(["hi"]), {f"extra/part{i}.xml": b"<x/>" for i in range(docparse.DOCX_MAX_ENTRIES)})
    with pytest.raises(docparse.ExtractError, match="too large"):
        docparse.docx_text(data)


def test_docx_document_xml_real_size_is_checked(monkeypatch):
    monkeypatch.setattr(docparse, "DOCX_MAX_DOCUMENT_XML", 2_000)
    with pytest.raises(docparse.ExtractError, match="too large"):
        docparse.docx_text(make_docx(["x" * 5_000]))


def test_docx_paragraph_bomb_is_refused_before_parsing(monkeypatch):
    monkeypatch.setattr(docparse, "DOCX_MAX_BLOCKS", 50)
    started = time.perf_counter()
    with pytest.raises(docparse.ExtractError, match="too large"):
        docparse.docx_text(make_docx(["p"] * 200))
    assert time.perf_counter() - started < 1


def test_docx_run_bomb_is_refused_before_parsing(monkeypatch):
    """One paragraph of 330k runs held a parse slot for the full 20 s in the review."""
    monkeypatch.setattr(docparse, "DOCX_MAX_RUNS", 50)
    document = docx.Document()
    paragraph = document.add_paragraph()
    for _ in range(80):
        paragraph.add_run("a")
    buf = io.BytesIO()
    document.save(buf)
    with pytest.raises(docparse.ExtractError, match="too large"):
        docparse.docx_text(buf.getvalue())


def test_normal_docx_still_reads():
    text = docparse.docx_text(make_docx(["Type designer in Busan."], [["Skill", "Level"], ["Lettering", "High"]]))
    assert "Type designer in Busan." in text and "Lettering | High" in text


# --------------------------------------------------------------------------- child process


async def test_isolated_parse_returns_text_and_friendly_errors():
    assert "brutalist posters" in await docparse.extract_isolated("pdf", make_pdf("Loves brutalist posters"), 20)
    assert "Seoul" in await docparse.extract_isolated("docx", make_docx(["Product designer in Seoul."]), 20)
    with pytest.raises(docparse.ExtractError, match="valid PDF"):
        await docparse.extract_isolated("pdf", b"not a pdf at all", 20)
    with pytest.raises(docparse.ExtractError, match="valid .docx"):
        await docparse.extract_isolated("docx", b"PK\x03\x04garbage", 20)


async def test_timeout_kills_the_child(monkeypatch):
    spawned = []
    real_spawn = docparse._spawn

    def spy():
        proc = real_spawn()
        spawned.append(proc)
        return proc

    monkeypatch.setattr(docparse, "_spawn", spy)
    with pytest.raises(TimeoutError):
        await docparse.extract_isolated("pdf", make_pdf("slow"), 0.01)  # can't even start in time
    assert spawned and spawned[0].poll() is not None  # killed and reaped, not left running


def test_memory_limit_turns_a_huge_allocation_into_memory_error():
    script = (
        "from app import docparse\n"
        "docparse._limit_memory(96 * 1024 * 1024)\n"
        "try:\n    bytearray(400 * 1024 * 1024)\nexcept MemoryError:\n    print('capped')\n"
        "else:\n    print('not capped')\n"
    )
    cmd, env = docparse._child_command()
    out = subprocess.run([cmd[0], "-c", script], capture_output=True, cwd=docparse._BACKEND_DIR, env=env, timeout=60)
    assert out.stdout.strip() == b"capped", out.stderr[-500:]


def test_child_caps_its_memory_before_reading_the_document(monkeypatch):
    events: list[str] = []

    class _Stdin:
        @property
        def buffer(self):
            events.append("read")
            return io.BytesIO(b"pdf\n" + make_pdf("Zine maker"))

    out = io.BytesIO()
    monkeypatch.setattr(docparse, "_limit_memory", lambda limit: events.append(f"limit {limit >> 20} MB"))
    monkeypatch.setenv(docparse.MEMORY_ENV, "256")
    monkeypatch.setattr(sys, "stdin", _Stdin())
    monkeypatch.setattr(sys, "stdout", type("Out", (), {"buffer": out})())
    docparse._child_main()
    assert events == ["limit 256 MB", "read"]
    assert json.loads(out.getvalue()) == {"ok": True, "text": "Zine maker"}


async def test_waiting_for_a_parse_slot_is_bounded(monkeypatch):
    monkeypatch.setattr(docparse, "QUEUE_WAIT_S", 0.05)
    slots = docparse._semaphore()
    for _ in range(docparse.PARSE_CONCURRENCY):  # every slot busy with slow hostile files
        await slots.acquire()
    try:
        started = time.perf_counter()
        with pytest.raises(docparse.ExtractError, match="busy"):
            await docparse.extract_isolated("pdf", make_pdf("hello"), 20, owner="5.6.7.8")
        assert time.perf_counter() - started < 1
    finally:
        for _ in range(docparse.PARSE_CONCURRENCY):
            slots.release()
    assert "hello" in await docparse.extract_isolated("pdf", make_pdf("hello"), 20, owner="5.6.7.8")


async def test_one_client_gets_one_parse_at_a_time(monkeypatch):
    monkeypatch.setattr(docparse, "QUEUE_WAIT_S", 0.05)
    mine = docparse._owner_semaphore("1.2.3.4")  # this client already has a parse running
    await mine.acquire()
    try:
        with pytest.raises(docparse.ExtractError, match="busy"):
            await docparse.extract_isolated("pdf", make_pdf("queued"), 20, owner="1.2.3.4")
        # other clients are not held up by it
        assert "other" in await docparse.extract_isolated("pdf", make_pdf("other"), 20, owner="9.9.9.9")
    finally:
        mine.release()


async def test_ingest_file_reports_a_timeout_as_a_friendly_item(monkeypatch):
    monkeypatch.setattr(context, "PARSE_TIMEOUT_S", 0.01)
    draft = await context._ingest_file(UploadedFile("cv.pdf", make_pdf("hello"), "application/pdf"))
    assert draft.status == "error" and draft.error == "This file took too long to read."


async def test_ingest_file_uses_the_isolated_parser(monkeypatch):
    calls = []

    async def fake(kind, data, timeout, owner=None):
        calls.append((kind, timeout, owner))
        return "Illustrator who prints risograph zines."

    monkeypatch.setattr(docparse, "extract_isolated", fake)
    draft = await context._ingest_file(UploadedFile("cv.docx", b"PK..", "application/octet-stream"), "1.2.3.4")
    assert draft.status == "ok" and "risograph" in draft.text
    assert calls == [("docx", context.PARSE_TIMEOUT_S, "1.2.3.4")]
